import csv
import io
import json
import re
from datetime import UTC, date, datetime, timedelta, time, timezone
from typing import Any, List, Sequence

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # pragma: no cover - optional for non-XLSX helper verification
    load_workbook = None
from sqlalchemy import and_, delete, func, inspect, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from models.scheduler import CourseScheduleEvent, OfficialCalendarEvent, PersonalTodo
from schemas.scheduler import (
    ScheduleAnalysisSummarySchema,
    ScheduleAggregateData,
    ScheduleAiContextSchema,
    ScheduleCandidateFocusBlockSchema,
    ScheduleCandidateRescheduleOptionSchema,
    ScheduleCandidateSplitOptionSchema,
    ScheduleCandidateSplitSegmentSchema,
    ScheduleConflictSchema,
    ScheduleEventSchema,
    SchedulePileupSchema,
    ScheduleSuggestionSchema,
)


TODO_ORDER_TIME = func.coalesce(PersonalTodo.start_time, PersonalTodo.deadline, PersonalTodo.end_time)
ROW_PREVIEW_LIMIT = 20
KEY_ALIASES = {
    "title": ["title", "name", "event", "event_name", "事项", "标题", "名称", "活动"],
    "description": ["description", "desc", "content", "notes", "remark", "备注", "说明", "内容"],
    "display_mode": ["display_mode", "displaymode", "mode", "render_mode", "show_as", "显示模式", "展示模式"],
    "source_year": ["source_year", "year", "calendar_year", "academic_year", "来源年份", "年份", "学年"],
    "event_type": ["event_type", "type", "category", "类别", "类型"],
    "location": ["location", "place", "venue", "room", "地点", "位置", "教室"],
    "course_name": ["course_name", "course", "course_title", "class_name", "课程", "课程名称", "课程名"],
    "course_code": ["course_code", "course_id", "code", "课程代码", "课程编号", "课程号"],
    "instructor": ["instructor", "teacher", "lecturer", "老师", "教师", "授课教师"],
    "date": ["date", "day", "日期"],
    "start_date": ["start_date", "date_start", "开始日期"],
    "end_date": ["end_date", "date_end", "结束日期"],
    "start_time": ["start_time", "start", "begin_time", "开始时间", "开始"],
    "end_time": ["end_time", "end", "finish_time", "结束时间", "结束"],
}
VALID_OFFICIAL_DISPLAY_MODES = {"background", "event"}
BLACKBOARD_IMPORT_SOURCE = "blackboard"
COURSE_SCHEDULE_IMPORT_SOURCE = "course_schedule_xlsx"
ASIA_SHANGHAI = timezone(timedelta(hours=8))
TERM_START_MONDAY = date(2026, 2, 23)
VALID_TERM_VALUES = {"spring", "fall"}
WEEKDAY_LABELS = {
    "星期一": 1, "周一": 1, "星期1": 1, "周1": 1, "1": 1,
    "星期二": 2, "周二": 2, "星期2": 2, "周2": 2, "2": 2,
    "星期三": 3, "周三": 3, "星期3": 3, "周3": 3, "3": 3,
    "星期四": 4, "周四": 4, "星期4": 4, "周4": 4, "4": 4,
    "星期五": 5, "周五": 5, "星期5": 5, "周5": 5, "5": 5,
    "星期六": 6, "周六": 6, "星期6": 6, "周6": 6, "6": 6,
    "星期日": 7, "星期天": 7, "周日": 7, "周天": 7, "星期7": 7, "周7": 7, "7": 7,
}
SECTION_TIMES = {
    "1-2": (time(8, 0), time(9, 50)),
    "3-4": (time(10, 20), time(12, 10)),
    "5-6": (time(14, 0), time(15, 50)),
    "7-8": (time(16, 20), time(18, 10)),
    "9-10": (time(19, 0), time(20, 50)),
}
SECTION_PATTERN = re.compile(r"第?\s*(\d+)\s*[-~—至]\s*(\d+)\s*节")
WEEK_INFO_PATTERN = re.compile(r"(\d+)\s*[-~—至]\s*(\d+)\s*周\s*(单周|双周)?")
TERM_PATTERN = re.compile(r"(第[一二三四五六七八九十]+学期|\d{4}\s*[-/]\s*\d{4}\s*学年.*学期|[春秋]季学期)")
INSTRUCTOR_HINT_PATTERN = re.compile(r"(教师|老师|授课教师|主讲)")
BRACKET_GROUP_PATTERN = re.compile(r"\[([^\]]+)\]")
LOCATION_PREFIX_PATTERN = re.compile(r"^(地点|教室|位置|上课地点|教学楼|教室地点)[:：]?\s*", re.IGNORECASE)
LOCATION_HINT_PATTERN = re.compile(r"(教室|地点|位置|楼|室|馆|厅|实验楼|实验室|报告厅|礼堂|Room|ROOM|线上|网课|校区|Lab|LAB|腾讯会议|会议号)")
CLASS_INFO_PATTERN = re.compile(
    r"(班级|行政班|教学班|自然班|分组|专业|学院|本科|硕士|博士|重修|辅修|必修|选修|公选|专选|专必|双语|语言|实验班|英才|"
    r"\d{2,}.*班|[(（]英[)）]|英语|英文班)"
)
ROOM_CODE_PATTERN = re.compile(r"([A-Za-z]{1,4}\s*[-_]?\s*\d{2,4}|\d+[A-Za-z]?\s*[-_]?\s*\d{2,4})")
ONLINE_LOCATION_PATTERN = re.compile(r"(线上|网课|腾讯会议|会议号|直播|ZOOM|Zoom|zoom)")
LATIN_NAME_TOKEN_PATTERN = re.compile(r"[A-Za-z][A-Za-z.'-]*")
SUGGESTION_DAY_START_HOUR = 7
SUGGESTION_DAY_END_HOUR = 22
DEFAULT_FOCUS_BLOCK_MINUTES = 120
MIN_SPLIT_SEGMENT_MINUTES = 60


async def create_todo(
    db: AsyncSession,
    user_id: int,
    title: str,
    content: str = None,
    status: str = "pending",
    priority: int = 2,
    deadline: datetime | None = None,
    start_time: datetime | None = None,
    end_time: datetime | None = None,
    import_source: str | None = None,
    import_uid: str | None = None,
) -> PersonalTodo:
    new_todo = PersonalTodo(
        user_id=user_id,
        title=title,
        content=content,
        status=status,
        priority=priority,
        deadline=deadline,
        start_time=start_time,
        end_time=end_time,
        import_source=import_source,
        import_uid=import_uid,
    )
    db.add(new_todo)
    await db.commit()
    await db.refresh(new_todo)
    return new_todo


async def get_user_todo(db: AsyncSession, user_id: int, todo_id: int) -> PersonalTodo | None:
    result = await db.execute(
        select(PersonalTodo).where(
            PersonalTodo.id == todo_id,
            PersonalTodo.user_id == user_id,
        )
    )
    return result.scalar_one_or_none()


async def list_user_todos(db: AsyncSession, user_id: int) -> List[PersonalTodo]:
    result = await db.execute(
        select(PersonalTodo)
        .where(PersonalTodo.user_id == user_id)
        .order_by(TODO_ORDER_TIME.asc(), PersonalTodo.id.asc())
    )
    return result.scalars().all()


async def list_user_todos_by_status(db: AsyncSession, user_id: int, status: str) -> List[PersonalTodo]:
    result = await db.execute(
        select(PersonalTodo)
        .where(
            PersonalTodo.user_id == user_id,
            PersonalTodo.status == status,
        )
        .order_by(TODO_ORDER_TIME.asc(), PersonalTodo.id.asc())
    )
    return result.scalars().all()


async def update_todo(
    db: AsyncSession,
    user_id: int,
    todo_id: int,
    updates: dict
) -> PersonalTodo | None:
    todo = await get_user_todo(db, user_id, todo_id)
    if not todo:
        return None

    for field, value in updates.items():
        setattr(todo, field, value)

    await db.commit()
    await db.refresh(todo)
    return todo


async def delete_todo(db: AsyncSession, user_id: int, todo_id: int) -> bool:
    result = await db.execute(
        delete(PersonalTodo).where(
            PersonalTodo.id == todo_id,
            PersonalTodo.user_id == user_id,
        )
    )
    await db.commit()
    return result.rowcount > 0


async def list_user_todos_by_import_source(
    db: AsyncSession,
    user_id: int,
    import_source: str,
) -> List[PersonalTodo]:
    result = await db.execute(
        select(PersonalTodo)
        .where(
            PersonalTodo.user_id == user_id,
            PersonalTodo.import_source == import_source,
        )
        .order_by(TODO_ORDER_TIME.asc(), PersonalTodo.id.asc())
    )
    return result.scalars().all()


def normalize_key(value: str) -> str:
    if not isinstance(value, str):
        return ""
    return re.sub(r"[\s_\-]+", "", value.strip().lower())


def normalize_cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in text.split("\n")]
    return "\n".join([line for line in lines if line])


def normalize_weekday_label(value: Any) -> int | None:
    text = normalize_cell_text(value).replace("（", "(").replace("）", ")")
    if not text:
        return None
    text = re.sub(r"\s+", "", text)
    for candidate, weekday in WEEKDAY_LABELS.items():
        if candidate in text:
            return weekday
    return None


def normalize_section_text(value: Any) -> str | None:
    text = normalize_cell_text(value)
    matched = SECTION_PATTERN.search(text)
    if not matched:
        return None
    return f"第{matched.group(1)}-{matched.group(2)}节"


def get_section_key(section_text: str | None) -> str | None:
    matched = SECTION_PATTERN.search(section_text or "")
    if not matched:
        return None
    return f"{matched.group(1)}-{matched.group(2)}"


def build_course_slot_datetimes(
    weekday: int,
    section_text: str,
    base_monday: date = TERM_START_MONDAY,
    week_index: int = 1,
) -> tuple[datetime, datetime]:
    section_key = get_section_key(section_text)
    start_clock, end_clock = SECTION_TIMES.get(section_key, (time(8, 0), time(9, 50)))
    target_date = base_monday + timedelta(days=(max(1, week_index) - 1) * 7 + max(0, weekday - 1))
    return datetime.combine(target_date, start_clock), datetime.combine(target_date, end_clock)


def get_monday_of_date(value: date) -> date:
    return value - timedelta(days=value.weekday())


def get_term_start_date(term: str, year: int) -> date:
    normalized_term = str(term or "").strip().lower()
    if normalized_term not in VALID_TERM_VALUES:
        raise ValueError("term 只允许 spring 或 fall")
    if not isinstance(year, int) or year < 2000 or year > 2100:
        raise ValueError("year 必须是有效的四位年份")
    if normalized_term == "spring":
        return date(year, 2, 25)
    return date(year, 9, 10)


def get_term_start_monday(term: str, year: int) -> date:
    return get_monday_of_date(get_term_start_date(term, year))


def build_term_label(term: str, year: int) -> str:
    normalized_term = str(term or "").strip().lower()
    if normalized_term == "spring":
        return f"{year}春季学期"
    if normalized_term == "fall":
        return f"{year}秋季学期"
    raise ValueError("term 只允许 spring 或 fall")


def extract_term_label(lines: list[str]) -> str | None:
    for line in lines:
        matched = TERM_PATTERN.search(line)
        if matched:
            return matched.group(1).replace(" ", "")
    return None


def normalize_week_rule(week_info: str | None) -> tuple[int | None, int | None, str]:
    text = clean_course_display_text(week_info)
    matched = WEEK_INFO_PATTERN.search(text)
    if not matched:
        matched = re.search(r"(\d+)\s*[-~—至]\s*(\d+)\s*(单周|双周)", text)
    if not matched:
        return None, None, "all"
    week_type = "all"
    suffix = matched.group(3) or ""
    if "单" in suffix:
        week_type = "odd"
    elif "双" in suffix:
        week_type = "even"
    return int(matched.group(1)), int(matched.group(2)), week_type


def extract_bracket_groups(value: Any) -> list[str]:
    text = "" if value is None else str(value)
    return [group.strip() for group in BRACKET_GROUP_PATTERN.findall(text) if str(group).strip()]


def clean_course_display_text(value: Any) -> str:
    text = normalize_cell_text(value).strip()
    if not text:
        return ""
    text = re.sub(r"^[\[\]【】]+|[\[\]【】]+$", "", text)
    text = LOCATION_PREFIX_PATTERN.sub("", text)
    text = re.sub(r"^(教师|老师|授课教师|主讲)[:：]?\s*", "", text)
    return text.strip()


def split_compound_text_tokens(value: Any) -> list[str]:
    text = normalize_cell_text(value)
    if not text:
        return []
    tokens: list[str] = []
    for part in re.split(r"[\n|｜]+", text):
        cleaned_part = part.strip()
        if not cleaned_part:
            continue
        bracket_tokens = extract_bracket_groups(cleaned_part)
        outside = clean_course_display_text(BRACKET_GROUP_PATTERN.sub(" ", cleaned_part))
        if outside:
            tokens.append(outside)
        tokens.extend(bracket_tokens)
    return [clean_course_display_text(token) for token in tokens if clean_course_display_text(token)]


def normalize_compare_text(value: Any) -> str:
    text = clean_course_display_text(value)
    if not text:
        return ""
    return re.sub(r"[\s\[\]【】()（）:：,，/／·.。_-]+", "", text).lower()


def is_duplicate_course_info(candidate: str, course_name: str) -> bool:
    candidate_key = normalize_compare_text(candidate)
    course_key = normalize_compare_text(course_name)
    if not candidate_key or not course_key:
        return False
    return candidate_key == course_key or candidate_key in course_key or course_key in candidate_key


def is_probable_location_text(candidate: str, course_name: str) -> bool:
    text = clean_course_display_text(candidate)
    if not text:
        return False
    if INSTRUCTOR_HINT_PATTERN.search(text):
        return False
    if WEEK_INFO_PATTERN.search(text) or SECTION_PATTERN.search(text) or TERM_PATTERN.search(text):
        return False
    if is_duplicate_course_info(text, course_name):
        return False
    has_location_hint = bool(
        LOCATION_HINT_PATTERN.search(text)
        or ROOM_CODE_PATTERN.search(text)
        or ONLINE_LOCATION_PATTERN.search(text)
        or LOCATION_PREFIX_PATTERN.search(candidate)
    )
    if CLASS_INFO_PATTERN.search(text) and not has_location_hint:
        return False
    return has_location_hint


def clean_instructor_token(value: Any) -> str | None:
    text = clean_course_display_text(value)
    if not text:
        return None
    text = re.sub(r"(第?\s*\d+\s*[-－—~至]\s*\d+\s*节.*)$", "", text).strip()
    text = re.sub(r"(\d+\s*[-~—至]\s*\d+\s*周\s*(单周|双周)?\s*)$", "", text).strip()
    text = LOCATION_PREFIX_PATTERN.sub("", text).strip(" ,，;；/|")
    return text or None


def looks_like_instructor_name(token: str) -> bool:
    text = clean_instructor_token(token)
    if not text:
        return False
    if re.fullmatch(r"[\u4e00-\u9fff·]{2,6}", text):
        return True
    latin_words = LATIN_NAME_TOKEN_PATTERN.findall(text)
    if len(latin_words) < 2 or len(latin_words) > 6:
        return False
    letters_only = re.sub(r"[^A-Za-z]", "", text)
    if len(letters_only) < 6 or len(letters_only) > 40:
        return False
    token_without_spaces = re.sub(r"\s+", "", text)
    latin_chars = re.sub(r"[^A-Za-z.' -]", "", text)
    if not latin_chars or len(latin_chars) < max(6, int(len(token_without_spaces) * 0.6)):
        return False
    return True


def is_probable_instructor_text(candidate: str, course_name: str) -> bool:
    text = clean_instructor_token(candidate)
    if not text:
        return False
    if is_duplicate_course_info(text, course_name):
        return False
    if WEEK_INFO_PATTERN.search(text) or SECTION_PATTERN.search(text) or TERM_PATTERN.search(text):
        return False
    if is_probable_location_text(text, course_name):
        return False
    if CLASS_INFO_PATTERN.search(text):
        return False
    if INSTRUCTOR_HINT_PATTERN.search(candidate):
        return True
    return looks_like_instructor_name(text)


def is_class_info_text(candidate: str) -> bool:
    text = clean_course_display_text(candidate)
    if not text:
        return False
    return bool(CLASS_INFO_PATTERN.search(text))


def normalize_week_info_text(value: Any) -> str | None:
    text = clean_course_display_text(value)
    if not text:
        return None
    matched = re.search(r"(\d+)\s*[-~—至]\s*(\d+)\s*周?\s*(单周|双周)?", text)
    if not matched:
        return None
    start_week = int(matched.group(1))
    end_week = int(matched.group(2))
    suffix = matched.group(3) or ""
    if suffix:
        return f"{start_week}-{end_week}{suffix}"
    return f"{start_week}-{end_week}周"


def normalize_section_label(value: Any) -> str | None:
    text = clean_course_display_text(value)
    if not text:
        return None
    matched = re.search(r"第?\s*(\d+)\s*[-－—~至]\s*(\d+)\s*节?", text)
    if not matched:
        return None
    return f"第{int(matched.group(1))}-{int(matched.group(2))}节"


def normalize_location_text(value: Any) -> str | None:
    text = clean_course_display_text(value)
    if not text:
        return None
    return LOCATION_PREFIX_PATTERN.sub("", text).strip() or None


def normalize_instructor_text(value: Any) -> str | None:
    text = clean_instructor_token(value)
    if not text:
        return None
    return re.sub(r"^(教师|老师|授课教师|主讲)[:：]?\s*", "", text).strip() or None


def build_course_schedule_dedupe_key(payload: dict[str, Any]) -> tuple[Any, ...]:
    def normalize_value(value: Any) -> Any:
        if isinstance(value, datetime):
            return value.isoformat(timespec="seconds")
        if value is None:
            return ""
        return normalize_compare_text(value)

    return (
        normalize_value(payload.get("user_id")),
        normalize_value(payload.get("title")),
        normalize_value(payload.get("course_name")),
        normalize_value(payload.get("course_code")),
        normalize_value(normalize_instructor_text(payload.get("instructor"))),
        normalize_value(normalize_location_text(payload.get("location"))),
        normalize_value(payload.get("weekday")),
        normalize_value(normalize_section_label(payload.get("section_text"))),
        normalize_value(payload.get("start_time")),
        normalize_value(payload.get("end_time")),
    )


def dedupe_course_schedule_payloads(payloads: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[tuple[Any, ...]] = set()
    deduped: list[dict[str, Any]] = []
    for payload in payloads:
        key = build_course_schedule_dedupe_key(payload)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(payload)
    return deduped


def parse_course_cell_text(
    cell_text: str,
    fallback_section_text: str,
    weekday: int,
    raw_text: str | None = None,
) -> dict[str, Any]:
    lines = [line.strip() for line in normalize_cell_text(cell_text).split("\n") if line.strip()]
    if not lines:
        raise ValueError("课程单元格为空")

    term_label = extract_term_label(lines)
    bracket_groups = extract_bracket_groups(raw_text if raw_text is not None else cell_text)

    title_candidates: list[str] = []
    trailing_tokens: list[str] = []
    for line in lines:
        if term_label and term_label in line:
            line = line.replace(term_label, " ")
        outside_brackets = clean_course_display_text(BRACKET_GROUP_PATTERN.sub(" ", line))
        if outside_brackets:
            title_candidates.append(outside_brackets)
        inline_groups = extract_bracket_groups(line)
        trailing_tokens.extend([clean_course_display_text(token) for token in inline_groups if clean_course_display_text(token)])

    course_name = next(
        (
            candidate for candidate in title_candidates
            if not normalize_week_info_text(candidate)
            and not normalize_section_label(candidate)
            and not is_class_info_text(candidate)
            and not INSTRUCTOR_HINT_PATTERN.search(candidate)
            and not is_probable_location_text(candidate, "")
        ),
        None,
    )
    if course_name is None:
        course_name = clean_course_display_text(lines[0])
    if not course_name:
        raise ValueError("未识别到课程名称")

    week_info = None
    section_text = normalize_section_label(fallback_section_text)
    location = None
    instructor = None
    class_infos: list[str] = []
    leftovers: list[str] = []

    classified_tokens: list[str] = []
    classified_tokens.extend([clean_course_display_text(token) for token in bracket_groups if clean_course_display_text(token)])
    classified_tokens.extend(trailing_tokens)
    for line in lines[1:] if len(lines) > 1 else []:
        classified_tokens.extend(split_compound_text_tokens(line))

    seen_tokens: set[str] = set()
    for token in classified_tokens:
        token = clean_course_display_text(token)
        token_key = normalize_compare_text(token)
        if not token or not token_key or token_key in seen_tokens:
            continue
        seen_tokens.add(token_key)

        normalized_week = normalize_week_info_text(token)
        if week_info is None and normalized_week is not None:
            week_info = normalized_week
            continue

        normalized_section = normalize_section_label(token)
        if section_text is None and normalized_section is not None:
            section_text = normalized_section
            continue

        if term_label and token == term_label:
            continue

        if is_duplicate_course_info(token, course_name):
            continue

        if location is None and is_probable_location_text(token, course_name):
            location = normalize_location_text(token)
            continue

        if instructor is None and is_probable_instructor_text(token, course_name):
            instructor = normalize_instructor_text(token)
            continue

        if is_class_info_text(token):
            class_infos.append(token)
            continue

        leftovers.append(token)

    if section_text is None:
        raise ValueError("未识别到节次")

    week_start, week_end, week_type = normalize_week_rule(week_info)
    start_time, end_time = build_course_slot_datetimes(weekday, section_text)

    description_parts = []
    if class_infos:
        description_parts.append("\n".join(class_infos))
    if leftovers:
        description_parts.append("\n".join(leftovers))
    if term_label:
        description_parts.append(f"学期: {term_label}")

    return {
        "title": course_name,
        "course_name": course_name,
        "instructor": instructor or None,
        "location": location or None,
        "week_info": week_info,
        "section_text": section_text,
        "weekday": weekday,
        "term_label": term_label,
        "raw_text": raw_text if raw_text is not None else str(cell_text),
        "week_start": week_start,
        "week_end": week_end,
        "week_type": week_type,
        "description": "\n".join(description_parts) if description_parts else None,
        "start_time": start_time,
        "end_time": end_time,
    }


def parse_course_schedule_xlsx(content: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    if not content:
        raise ValueError("XLSX 文件内容为空")
    if load_workbook is None:
        raise ValueError("缺少 openpyxl 依赖，无法解析 XLSX 课程表")

    workbook = load_workbook(filename=io.BytesIO(content), data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("XLSX 首个工作表为空")

    header_row_index = None
    weekday_columns: dict[int, int] = {}
    for row_index, row in enumerate(rows):
        current_map = {}
        for col_index, value in enumerate(row):
            weekday = normalize_weekday_label(value)
            if weekday is not None:
                current_map[col_index] = weekday
        if len(current_map) >= 3:
            header_row_index = row_index
            weekday_columns = current_map
            break

    if header_row_index is None:
        raise ValueError("未识别到课程表星期表头行")

    payloads: list[dict[str, Any]] = []
    skipped_rows: list[str] = []
    for row_index in range(header_row_index + 1, len(rows)):
        row = rows[row_index]
        section_text = next((normalize_section_text(value) for value in row if normalize_section_text(value)), None)
        if not section_text:
            continue

        for col_index, weekday in weekday_columns.items():
            if col_index >= len(row):
                continue
            raw_cell_value = row[col_index]
            raw_text = "" if raw_cell_value is None else str(raw_cell_value)
            normalized_text = normalize_cell_text(raw_cell_value)
            if not normalized_text:
                continue
            try:
                payload = parse_course_cell_text(
                    normalized_text,
                    section_text,
                    weekday,
                    raw_text=raw_text,
                )
                payloads.append(payload)
            except ValueError as error:
                skipped_rows.append(f"第 {row_index + 1} 行 星期{weekday} 解析失败: {error}")

    if not payloads:
        raise ValueError("未从 XLSX 课程表中解析到任何课程")

    return dedupe_course_schedule_payloads(payloads), skipped_rows


def should_expand_week(week_index: int, week_type: str | None) -> bool:
    normalized_type = (week_type or "all").strip().lower()
    if normalized_type == "odd":
        return week_index % 2 == 1
    if normalized_type == "even":
        return week_index % 2 == 0
    return True


def expand_course_schedule_payloads(
    payloads: list[dict[str, Any]],
    term_start_monday: date = TERM_START_MONDAY,
) -> list[dict[str, Any]]:
    expanded_events: list[dict[str, Any]] = []

    for payload in payloads:
        weekday = int(payload.get("weekday") or 0)
        section_text = payload.get("section_text")
        week_start = payload.get("week_start")
        week_end = payload.get("week_end")
        week_type = payload.get("week_type") or "all"

        if weekday <= 0 or not section_text:
            expanded_events.append({**payload})
            continue

        if week_start is None or week_end is None:
            expanded_events.append({**payload})
            continue

        if week_end < week_start:
            continue

        for week_index in range(week_start, week_end + 1):
            if not should_expand_week(week_index, week_type):
                continue

            start_time, end_time = build_course_slot_datetimes(
                weekday=weekday,
                section_text=section_text,
                base_monday=term_start_monday,
                week_index=week_index,
            )
            expanded_events.append(
                {
                    **payload,
                    "start_time": start_time,
                    "end_time": end_time,
                }
            )

    return dedupe_course_schedule_payloads(expanded_events)


async def ensure_course_schedule_schema(db: AsyncSession):
    def sync_ensure(sync_session):
        conn = sync_session.connection()

        existing_columns = {
            column["name"]
            for column in inspect(conn).get_columns(CourseScheduleEvent.__tablename__)
        }

        required_columns = {
            "week_info": "ALTER TABLE course_schedule_event ADD COLUMN week_info VARCHAR(255) NULL",
            "section_text": "ALTER TABLE course_schedule_event ADD COLUMN section_text VARCHAR(100) NULL",
            "weekday": "ALTER TABLE course_schedule_event ADD COLUMN weekday INT NULL",
            "term_label": "ALTER TABLE course_schedule_event ADD COLUMN term_label VARCHAR(255) NULL",
            "raw_text": "ALTER TABLE course_schedule_event ADD COLUMN raw_text TEXT NULL",
            "week_start": "ALTER TABLE course_schedule_event ADD COLUMN week_start INT NULL",
            "week_end": "ALTER TABLE course_schedule_event ADD COLUMN week_end INT NULL",
            "week_type": "ALTER TABLE course_schedule_event ADD COLUMN week_type VARCHAR(20) NULL",
            "import_source": "ALTER TABLE course_schedule_event ADD COLUMN import_source VARCHAR(50) NULL",
            "import_uid": "ALTER TABLE course_schedule_event ADD COLUMN import_uid VARCHAR(255) NULL",
        }

        for column_name, sql in required_columns.items():
            if column_name not in existing_columns:
                conn.exec_driver_sql(sql)

    await db.run_sync(sync_ensure)
    await db.commit()

def get_row_value(row: dict[str, Any], field: str) -> Any:
    aliases = KEY_ALIASES.get(field, [field])
    for key, value in row.items():
        if normalize_key(str(key)) in {normalize_key(alias) for alias in aliases}:
            return value
    return None


def parse_date_value(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()

    text = str(value).strip()
    if not text:
        return None

    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    return None


def parse_time_value(value: Any) -> time | None:
    if value is None:
        return None
    if isinstance(value, time):
        return value.replace(tzinfo=None)
    if isinstance(value, datetime):
        return value.time().replace(tzinfo=None)

    text = str(value).strip()
    if not text:
        return None

    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue

    return None


def parse_datetime_value(value: Any) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time.min)

    text = str(value).strip()
    if not text:
        return None

    text = text.replace("Z", "+00:00")
    for candidate in (text, text.replace("/", "-"), text.replace(".", "-")):
        try:
            parsed = datetime.fromisoformat(candidate)
            return parsed.replace(tzinfo=None)
        except ValueError:
            pass

    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y.%m.%d %H:%M:%S",
        "%Y.%m.%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
        "%Y%m%d %H:%M",
        "%Y%m%d",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue

    return None


def parse_ics_datetime_value(value: str, tzid: str | None = None) -> datetime | None:
    text = (value or "").strip()
    if not text:
        return None

    if text.endswith("Z"):
        for fmt in ("%Y%m%dT%H%M%SZ",):
            try:
                parsed = datetime.strptime(text, fmt).replace(tzinfo=UTC)
                return parsed.astimezone(ASIA_SHANGHAI).replace(tzinfo=None)
            except ValueError:
                continue

    for fmt in ("%Y%m%dT%H%M%S", "%Y%m%dT%H%M"):
        try:
            parsed = datetime.strptime(text, fmt)
            if tzid == "Asia/Shanghai":
                return parsed
            return parsed
        except ValueError:
            continue

    if len(text) == 8:
        try:
            parsed_date = datetime.strptime(text, "%Y%m%d").date()
            return datetime.combine(parsed_date, time.min)
        except ValueError:
            return None

    return None


def parse_int_value(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value

    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d{4}", text):
        return int(text)
    return None


def normalize_display_mode(value: Any) -> str | None:
    if value is None:
        return "event"

    normalized = normalize_key(str(value))
    mapping = {
        "background": "background",
        "bg": "background",
        "marker": "background",
        "layer": "background",
        "context": "background",
        "背景": "background",
        "背景层": "background",
        "event": "event",
        "activity": "event",
        "活动": "event",
        "事件": "event",
    }
    result = mapping.get(normalized)
    if result in VALID_OFFICIAL_DISPLAY_MODES:
        return result
    return None


def combine_datetime_parts(date_value: Any, time_value: Any) -> datetime | None:
    parsed_date = parse_date_value(date_value)
    if parsed_date is None:
        return None
    parsed_time = parse_time_value(time_value) or time.min
    return datetime.combine(parsed_date, parsed_time)


def resolve_row_time_range(row: dict[str, Any]) -> tuple[datetime | None, datetime | None]:
    start_time_value = get_row_value(row, "start_time")
    end_time_value = get_row_value(row, "end_time")
    direct_start_dt = parse_datetime_value(start_time_value)
    direct_end_dt = parse_datetime_value(end_time_value)
    start_dt = direct_start_dt
    end_dt = direct_end_dt
    shared_date = get_row_value(row, "date")
    start_date = get_row_value(row, "start_date") or shared_date
    end_date = get_row_value(row, "end_date") or shared_date

    if start_dt is None:
        start_dt = combine_datetime_parts(start_date, start_time_value)
    if end_dt is None:
        end_dt = combine_datetime_parts(end_date, end_time_value)

    if direct_start_dt is None and start_dt is not None and parse_time_value(start_time_value) is None and parse_date_value(start_date) is not None:
        start_dt = datetime.combine(start_dt.date(), time.min)
    if direct_end_dt is None and end_dt is not None and parse_time_value(end_time_value) is None and parse_date_value(end_date) is not None:
        end_dt = datetime.combine(end_dt.date(), time(23, 59, 59))

    if start_dt is None and end_dt is None and shared_date is not None:
        day = parse_date_value(shared_date)
        if day is not None:
            start_dt = datetime.combine(day, time.min)
            end_dt = datetime.combine(day, time(23, 59, 59))

    if start_dt is not None and end_dt is None:
        parsed_end_date = parse_date_value(end_date)
        if parsed_end_date is not None and parsed_end_date != start_dt.date():
            end_dt = datetime.combine(parsed_end_date, time(23, 59, 59))

    return start_dt, end_dt


def parse_rows_from_text(content: str, content_type: str | None) -> list[dict[str, Any]]:
    text = (content or "").strip()
    if not text:
        raise ValueError("导入内容不能为空")

    normalized_type = (content_type or "auto").strip().lower()

    if normalized_type == "json" or (normalized_type == "auto" and text[:1] in "[{"):
        data = json.loads(text)
        if isinstance(data, list):
            return [item for item in data if isinstance(item, dict)]
        if isinstance(data, dict):
            for key in ("rows", "data", "items", "events"):
                candidate = data.get(key)
                if isinstance(candidate, list):
                    return [item for item in candidate if isinstance(item, dict)]
        raise ValueError("JSON 内容必须是对象数组，或包含 rows/data/items/events 数组")

    delimiter = ","
    if normalized_type in {"tsv", "table"}:
        delimiter = "\t"
    elif normalized_type == "csv":
        delimiter = ","
    elif "|" in text and "\n" in text:
        delimiter = "|"
    elif "\t" in text:
        delimiter = "\t"

    lines = [line for line in text.splitlines() if line.strip()]
    if delimiter == "|" and len(lines) >= 2:
        filtered_lines = []
        for line in lines:
            stripped = line.strip()
            if re.fullmatch(r"\|?[\s:\-]+\|[\s:\-|]*", stripped):
                continue
            filtered_lines.append(stripped.strip("|"))
        lines = filtered_lines

    reader = csv.DictReader(io.StringIO("\n".join(lines)), delimiter=delimiter)
    if not reader.fieldnames:
        raise ValueError("表格内容缺少表头")

    rows = []
    for row in reader:
        clean_row = {}
        for key, value in row.items():
            if key is None:
                continue
            clean_row[str(key).strip()] = value.strip() if isinstance(value, str) else value
        if any(value not in (None, "") for value in clean_row.values()):
            rows.append(clean_row)
    return rows


def parse_import_rows(rows: list[dict[str, Any]] | None, content: str | None, content_type: str | None) -> list[dict[str, Any]]:
    if rows:
        return [row for row in rows if isinstance(row, dict)]
    if content is not None:
        return parse_rows_from_text(content, content_type)
    raise ValueError("请提供 rows 或 content")


def unfold_ics_lines(content: str) -> list[str]:
    raw_lines = content.splitlines()
    unfolded = []
    for line in raw_lines:
        if line.startswith((" ", "\t")) and unfolded:
            unfolded[-1] += line[1:]
        else:
            unfolded.append(line.rstrip("\r"))
    return unfolded


def parse_ics_property_line(line: str) -> tuple[str, dict[str, str], str]:
    if ":" not in line:
        return line.strip().upper(), {}, ""

    name_part, value = line.split(":", 1)
    parts = name_part.split(";")
    name = parts[0].strip().upper()
    params = {}
    for item in parts[1:]:
        if "=" not in item:
            continue
        key, param_value = item.split("=", 1)
        params[key.strip().upper()] = param_value.strip()
    return name, params, value.strip()


def parse_blackboard_ics_events(content: str) -> tuple[list[dict[str, Any]], list[str]]:
    lines = unfold_ics_lines(content or "")
    events = []
    skipped_rows = []
    current_event: dict[str, Any] | None = None

    for line in lines:
        stripped = line.strip()
        if stripped == "BEGIN:VEVENT":
            current_event = {"raw": []}
            continue
        if stripped == "END:VEVENT":
            if current_event is not None:
                if current_event.get("unsupported_reason"):
                    skipped_rows.append(current_event["unsupported_reason"])
                else:
                    events.append(current_event)
            current_event = None
            continue
        if current_event is None:
            continue

        name, params, value = parse_ics_property_line(stripped)
        current_event["raw"].append(stripped)

        if name == "RRULE":
            current_event["unsupported_reason"] = "发现 RRULE 重复规则，当前版本已跳过该 Blackboard 事件"
            continue

        if name in {"DTSTART", "DTEND"}:
            tzid = params.get("TZID")
            parsed_dt = parse_ics_datetime_value(value, tzid=tzid)
            current_event[name] = parsed_dt
            current_event[f"{name}_TZID"] = tzid
            continue

        if name in {"SUMMARY", "UID", "DESCRIPTION", "LOCATION"}:
            current_event[name] = value.replace("\\n", "\n").replace("\\,", ",")

    return events, skipped_rows


def build_blackboard_todo_payloads(user_id: int, content: str) -> tuple[list[dict[str, Any]], list[str]]:
    events, skipped_rows = parse_blackboard_ics_events(content)
    payloads = []

    for index, event in enumerate(events, start=1):
        summary = (event.get("SUMMARY") or "").strip()
        deadline = event.get("DTSTART")
        uid = (event.get("UID") or "").strip()
        if not summary or deadline is None:
            skipped_rows.append(f"第 {index} 个 Blackboard VEVENT 缺少 SUMMARY 或 DTSTART")
            continue

        description = (event.get("DESCRIPTION") or "").strip()
        location = (event.get("LOCATION") or "").strip()
        content_parts = []
        if description:
            content_parts.append(description)
        if location:
            content_parts.append(f"Location: {location}")
        if uid:
            content_parts.append(f"[Blackboard UID] {uid}")

        payloads.append(
            {
                "user_id": user_id,
                "title": summary,
                "content": "\n\n".join(content_parts) if content_parts else None,
                "status": "pending",
                "priority": 2,
                "deadline": deadline,
                "start_time": None,
                "end_time": None,
                "import_source": BLACKBOARD_IMPORT_SOURCE,
                "import_uid": uid or None,
            }
        )

    return payloads, skipped_rows


def build_official_calendar_payloads(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    payloads = []
    skipped_rows = []

    for index, row in enumerate(rows, start=1):
        title = get_row_value(row, "title")
        start_time, end_time = resolve_row_time_range(row)
        display_mode = normalize_display_mode(get_row_value(row, "display_mode"))
        source_year = parse_int_value(get_row_value(row, "source_year"))
        if not title or start_time is None or end_time is None or end_time < start_time:
            skipped_rows.append(f"第 {index} 行缺少合法的标题或开始/结束时间")
            continue
        if display_mode is None:
            skipped_rows.append(f"第 {index} 行 display_mode 不合法，只允许 background 或 event")
            continue
        if source_year is None:
            skipped_rows.append(f"第 {index} 行缺少合法的 source_year")
            continue

        payloads.append(
            {
                "title": str(title).strip(),
                "description": get_row_value(row, "description"),
                "display_mode": display_mode,
                "source_year": source_year,
                "event_type": str(get_row_value(row, "event_type") or "event").strip(),
                "start_time": start_time,
                "end_time": end_time,
                "location": get_row_value(row, "location"),
            }
        )

    return payloads, skipped_rows


def build_course_schedule_payloads(user_id: int, rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    payloads = []
    skipped_rows = []

    for index, row in enumerate(rows, start=1):
        course_name = get_row_value(row, "course_name") or get_row_value(row, "title")
        title = get_row_value(row, "title") or course_name
        start_time, end_time = resolve_row_time_range(row)
        if not course_name or not title or start_time is None or end_time is None or end_time < start_time:
            skipped_rows.append(f"第 {index} 行缺少合法的课程名称或开始/结束时间")
            continue

        payloads.append(
            {
                "user_id": user_id,
                "title": str(title).strip(),
                "course_name": str(course_name).strip(),
                "course_code": get_row_value(row, "course_code"),
                "instructor": get_row_value(row, "instructor"),
                "description": get_row_value(row, "description"),
                "week_info": get_row_value(row, "week_info"),
                "section_text": get_row_value(row, "section_text"),
                "weekday": parse_int_value(get_row_value(row, "weekday")),
                "term_label": get_row_value(row, "term_label"),
                "raw_text": get_row_value(row, "raw_text"),
                "week_start": parse_int_value(get_row_value(row, "week_start")),
                "week_end": parse_int_value(get_row_value(row, "week_end")),
                "week_type": get_row_value(row, "week_type") or "all",
                "import_source": get_row_value(row, "import_source"),
                "import_uid": get_row_value(row, "import_uid"),
                "start_time": start_time,
                "end_time": end_time,
                "location": get_row_value(row, "location"),
            }
        )

    return dedupe_course_schedule_payloads(payloads), skipped_rows


async def replace_official_calendar_events(
    db: AsyncSession,
    events: Sequence[dict[str, Any]],
    replace_existing: bool = True,
) -> List[OfficialCalendarEvent]:
    if replace_existing:
        await db.execute(delete(OfficialCalendarEvent))

    created = [OfficialCalendarEvent(**payload) for payload in events]
    db.add_all(created)
    await db.commit()
    for item in created:
        await db.refresh(item)
    return created


async def list_official_calendar_events(
    db: AsyncSession,
    source_year: int | None = None,
    display_mode: str | None = None,
    start_time: datetime | None = None,
    end_time: datetime | None = None,
) -> List[OfficialCalendarEvent]:
    stmt = select(OfficialCalendarEvent)
    if source_year is not None:
        stmt = stmt.where(OfficialCalendarEvent.source_year == source_year)
    if display_mode is not None:
        stmt = stmt.where(OfficialCalendarEvent.display_mode == display_mode)
    if start_time is not None:
        stmt = stmt.where(OfficialCalendarEvent.end_time >= start_time)
    if end_time is not None:
        stmt = stmt.where(OfficialCalendarEvent.start_time <= end_time)
    stmt = stmt.order_by(OfficialCalendarEvent.start_time.asc(), OfficialCalendarEvent.id.asc())
    result = await db.execute(stmt)
    return result.scalars().all()


async def replace_course_schedule_events(
    db: AsyncSession,
    user_id: int,
    events: Sequence[dict[str, Any]],
    replace_existing: bool = True,
) -> List[CourseScheduleEvent]:
    if replace_existing:
        await db.execute(delete(CourseScheduleEvent).where(CourseScheduleEvent.user_id == user_id))

    deduped_events = dedupe_course_schedule_payloads(events)
    created = [CourseScheduleEvent(**payload) for payload in deduped_events]
    db.add_all(created)
    await db.commit()
    for item in created:
        await db.refresh(item)
    return created


async def replace_blackboard_todos(
    db: AsyncSession,
    user_id: int,
    todos: Sequence[dict[str, Any]],
    replace_existing: bool = True,
) -> List[PersonalTodo]:
    if replace_existing:
        await db.execute(
            delete(PersonalTodo).where(
                PersonalTodo.user_id == user_id,
                PersonalTodo.import_source == BLACKBOARD_IMPORT_SOURCE,
            )
        )

    created = [PersonalTodo(**payload) for payload in todos]
    db.add_all(created)
    await db.commit()
    for item in created:
        await db.refresh(item)
    return created


async def list_course_schedule_events(
    db: AsyncSession,
    user_id: int,
    start_time: datetime | None = None,
    end_time: datetime | None = None,
) -> List[CourseScheduleEvent]:
    stmt = select(CourseScheduleEvent).where(CourseScheduleEvent.user_id == user_id)
    if start_time is not None:
        stmt = stmt.where(CourseScheduleEvent.end_time >= start_time)
    if end_time is not None:
        stmt = stmt.where(CourseScheduleEvent.start_time <= end_time)
    stmt = stmt.order_by(CourseScheduleEvent.start_time.asc(), CourseScheduleEvent.id.asc())
    result = await db.execute(stmt)
    return result.scalars().all()


async def list_blackboard_todos(
    db: AsyncSession,
    user_id: int,
) -> List[PersonalTodo]:
    result = await db.execute(
        select(PersonalTodo)
        .where(
            PersonalTodo.user_id == user_id,
            PersonalTodo.import_source == BLACKBOARD_IMPORT_SOURCE,
        )
        .order_by(TODO_ORDER_TIME.asc(), PersonalTodo.id.asc())
    )
    return result.scalars().all()


def _build_schedule_primary_time(event: ScheduleEventSchema) -> datetime | None:
    return event.start_time or event.deadline


def _has_explicit_range(event: ScheduleEventSchema) -> bool:
    return event.start_time is not None and event.end_time is not None


def _is_pending_deadline_todo(event: ScheduleEventSchema) -> bool:
    return (
        event.source == "personal_todo"
        and event.deadline is not None
        and not _has_explicit_range(event)
        and event.status not in {"done", "cancelled"}
    )


def _calculate_risk_level(
    conflicts: List[ScheduleConflictSchema],
    pileups: List[SchedulePileupSchema],
    suggestions: List[ScheduleSuggestionSchema],
) -> str:
    if any(item.level == "high" for item in conflicts) or any(item.level == "high" for item in pileups):
        return "high"
    if conflicts or pileups or any(item.level in {"medium", "high"} for item in suggestions):
        return "medium"
    return "low"


def _build_schedule_sort_key(event: ScheduleEventSchema) -> tuple:
    """
    Timeline ordering rule:
    1. Scheduled items are ordered by start_time ascending.
    2. Deadline-only items are ordered by deadline ascending.
    3. Items with the same primary time place scheduled items before deadline-only items.
    4. Remaining ties are broken by end_time, then source, then source_id for deterministic output.
    """
    primary_time = _build_schedule_primary_time(event) or datetime.max
    is_deadline_only = event.start_time is None and event.deadline is not None
    end_marker = event.end_time or event.deadline or datetime.max
    return (
        primary_time,
        1 if is_deadline_only else 0,
        end_marker,
        event.source,
        event.source_id,
    )


async def load_personal_todos_by_window(
    db: AsyncSession,
    user_id: int,
    start_time: datetime | None = None,
    end_time: datetime | None = None,
    include_blackboard: bool = True,
    include_completed: bool = True,
) -> List[PersonalTodo]:
    stmt = select(PersonalTodo).where(PersonalTodo.user_id == user_id)

    if not include_blackboard:
        stmt = stmt.where(
            or_(
                PersonalTodo.import_source.is_(None),
                PersonalTodo.import_source != BLACKBOARD_IMPORT_SOURCE,
            )
        )

    if not include_completed:
        stmt = stmt.where(PersonalTodo.status != "done")

    scheduled_overlap_clauses = []
    if start_time is not None:
        scheduled_overlap_clauses.append(PersonalTodo.end_time >= start_time)
    if end_time is not None:
        scheduled_overlap_clauses.append(PersonalTodo.start_time <= end_time)

    scheduled_window_clause = and_(
        PersonalTodo.start_time.is_not(None),
        PersonalTodo.end_time.is_not(None),
        *scheduled_overlap_clauses,
    )

    deadline_window_clauses = [PersonalTodo.deadline.is_not(None)]
    if start_time is not None:
        deadline_window_clauses.append(PersonalTodo.deadline >= start_time)
    if end_time is not None:
        deadline_window_clauses.append(PersonalTodo.deadline <= end_time)
    deadline_window_clause = and_(*deadline_window_clauses)

    if start_time is not None or end_time is not None:
        stmt = stmt.where(or_(scheduled_window_clause, deadline_window_clause))

    stmt = stmt.order_by(TODO_ORDER_TIME.asc(), PersonalTodo.id.asc())
    result = await db.execute(stmt)
    return result.scalars().all()


async def load_official_calendar_events_by_window(
    db: AsyncSession,
    start_time: datetime | None = None,
    end_time: datetime | None = None,
    include_background_markers: bool = True,
) -> List[OfficialCalendarEvent]:
    display_mode = None if include_background_markers else "event"
    return await list_official_calendar_events(
        db=db,
        display_mode=display_mode,
        start_time=start_time,
        end_time=end_time,
    )


def normalize_personal_todo_to_schedule_event(todo: PersonalTodo) -> ScheduleEventSchema:
    """
    Personal todos are user-managed items, so scheduled personal blocks are movable by default.
    In this backend, `is_flexible` means the rule layer may generate legal adjustment options for the event.
    """
    is_scheduled = todo.start_time is not None and todo.end_time is not None
    source_meta: dict[str, Any] = {
        "import_source": todo.import_source,
        "import_uid": todo.import_uid,
    }
    if todo.import_source == BLACKBOARD_IMPORT_SOURCE:
        source_meta["origin_domain"] = "blackboard"

    return ScheduleEventSchema(
        source="personal_todo",
        source_type="personal_todo",
        source_id=todo.id,
        user_id=todo.user_id,
        kind="scheduled_todo" if is_scheduled else "deadline_todo",
        title=todo.title,
        description=todo.content,
        start_time=todo.start_time,
        end_time=todo.end_time,
        deadline=todo.deadline,
        status=todo.status,
        priority=todo.priority,
        display_mode="event",
        is_blocking=is_scheduled,
        is_flexible=True,
        source_meta=source_meta,
    )


def normalize_official_calendar_event_to_schedule_event(
    event: OfficialCalendarEvent,
) -> ScheduleEventSchema:
    is_background = event.display_mode == "background"
    return ScheduleEventSchema(
        source="official_calendar_event",
        source_type="official_calendar_event",
        source_id=event.id,
        kind="background_marker" if is_background else "official_event",
        title=event.title,
        description=event.description,
        start_time=event.start_time,
        end_time=event.end_time,
        location=event.location,
        display_mode=event.display_mode,
        is_blocking=not is_background,
        is_flexible=False,
        source_meta={
            "source_year": event.source_year,
            "event_type": event.event_type,
            "display_mode": event.display_mode,
        },
    )


def normalize_course_schedule_event_to_schedule_event(
    event: CourseScheduleEvent,
) -> ScheduleEventSchema:
    return ScheduleEventSchema(
        source="course_schedule_event",
        source_type="course_schedule_event",
        source_id=event.id,
        user_id=event.user_id,
        kind="course_event",
        title=event.title,
        description=event.description,
        start_time=event.start_time,
        end_time=event.end_time,
        deadline=None,
        status=None,
        location=event.location,
        display_mode="event",
        is_blocking=True,
        is_flexible=False,
        source_meta={
            "course_name": event.course_name,
            "course_code": event.course_code,
            "instructor": event.instructor,
            "term_label": event.term_label,
            "weekday": event.weekday,
            "section_text": event.section_text,
            "week_info": event.week_info,
            "week_start": event.week_start,
            "week_end": event.week_end,
            "week_type": event.week_type,
            "raw_text": event.raw_text,
            "import_source": event.import_source,
            "import_uid": event.import_uid,
        },
    )


def merge_schedule_timeline(
    personal_todos: List[PersonalTodo],
    official_events: List[OfficialCalendarEvent],
    course_events: List[CourseScheduleEvent] | None = None,
) -> List[ScheduleEventSchema]:
    timeline = [normalize_personal_todo_to_schedule_event(todo) for todo in personal_todos]
    timeline.extend(normalize_official_calendar_event_to_schedule_event(event) for event in official_events)
    timeline.extend(normalize_course_schedule_event_to_schedule_event(event) for event in (course_events or []))
    return sorted(timeline, key=_build_schedule_sort_key)


async def build_schedule_timeline(
    db: AsyncSession,
    user_id: int,
    start_time: datetime | None = None,
    end_time: datetime | None = None,
    include_blackboard: bool = True,
    include_background_markers: bool = True,
    include_completed: bool = True,
    include_course_schedule: bool = True,
) -> List[ScheduleEventSchema]:
    personal_todos = await load_personal_todos_by_window(
        db=db,
        user_id=user_id,
        start_time=start_time,
        end_time=end_time,
        include_blackboard=include_blackboard,
        include_completed=include_completed,
    )
    official_events = await load_official_calendar_events_by_window(
        db=db,
        start_time=start_time,
        end_time=end_time,
        include_background_markers=include_background_markers,
    )
    course_events = []
    if include_course_schedule:
        course_events = await list_course_schedule_events(
            db=db,
            user_id=user_id,
            start_time=start_time,
            end_time=end_time,
        )
    return merge_schedule_timeline(personal_todos, official_events, course_events)


def detect_time_overlap_conflicts(
    schedule_events: List[ScheduleEventSchema],
) -> List[ScheduleConflictSchema]:
    """
    V1 overlap analysis only considers blocking items with explicit start/end times.
    Deadline-only tasks and background markers are excluded because no reliable duration is stored.
    """
    ranged_events = [
        event for event in schedule_events
        if event.is_blocking and _has_explicit_range(event)
    ]
    conflicts: list[ScheduleConflictSchema] = []

    for index, left_event in enumerate(ranged_events):
        for right_event in ranged_events[index + 1:]:
            if not (left_event.start_time < right_event.end_time and right_event.start_time < left_event.end_time):
                continue
            overlap_start = max(left_event.start_time, right_event.start_time)
            overlap_end = min(left_event.end_time, right_event.end_time)
            conflicts.append(
                ScheduleConflictSchema(
                    left_event=left_event,
                    right_event=right_event,
                    level="high",
                    reason="两个阻塞型时间段发生重叠",
                    conflict_type="time_overlap",
                    overlap_start_time=overlap_start,
                    overlap_end_time=overlap_end,
                    is_blocking=True,
                )
            )

    return conflicts


def detect_deadline_clusters(
    schedule_events: List[ScheduleEventSchema],
    same_day_threshold: int = 3,
    short_window_hours: int = 6,
    short_window_threshold: int = 2,
) -> List[SchedulePileupSchema]:
    """
    V1 deadline clustering only looks at pending deadline-only personal todos.
    It groups by calendar day and also flags dense short windows without inventing durations.
    """
    deadline_events = sorted(
        [event for event in schedule_events if _is_pending_deadline_todo(event)],
        key=lambda item: (item.deadline, item.source_id),
    )
    pileups: list[SchedulePileupSchema] = []

    events_by_day: dict[date, list[ScheduleEventSchema]] = {}
    for event in deadline_events:
        events_by_day.setdefault(event.deadline.date(), []).append(event)

    for day, events in sorted(events_by_day.items(), key=lambda item: item[0]):
        if len(events) >= same_day_threshold:
            pileups.append(
                SchedulePileupSchema(
                    start_time=min(item.deadline for item in events),
                    end_time=max(item.deadline for item in events),
                    event_count=len(events),
                    level="high" if len(events) >= same_day_threshold + 1 else "medium",
                    reason="同一天待完成截止事项过多",
                    events=events,
                )
            )

    for index, left_event in enumerate(deadline_events):
        window_events = [left_event]
        for right_event in deadline_events[index + 1:]:
            hours_apart = (right_event.deadline - left_event.deadline).total_seconds() / 3600
            if hours_apart > short_window_hours:
                break
            window_events.append(right_event)

        if len(window_events) < short_window_threshold:
            continue

        if any(
            pileup.start_time == window_events[0].deadline
            and pileup.end_time == window_events[-1].deadline
            and pileup.reason == "短时间窗口内待完成截止事项密集"
            for pileup in pileups
        ):
            continue

        if len(window_events) >= short_window_threshold:
            pileups.append(
                SchedulePileupSchema(
                    start_time=window_events[0].deadline,
                    end_time=window_events[-1].deadline,
                    event_count=len(window_events),
                    level="medium",
                    reason="短时间窗口内待完成截止事项密集",
                    events=window_events,
                )
            )

    return pileups


def _is_background_pressure_marker(event: ScheduleEventSchema) -> bool:
    return (
        event.source == "official_calendar_event"
        and event.kind == "background_marker"
        and event.display_mode == "background"
    )


def _time_window_overlaps(
    event: ScheduleEventSchema,
    start_time: datetime | None,
    end_time: datetime | None,
) -> bool:
    if not _has_explicit_range(event) or start_time is None or end_time is None:
        return False
    return event.start_time < end_time and start_time < event.end_time


def _build_event_key(event: ScheduleEventSchema) -> tuple[str, int]:
    return (event.source, event.source_id)


def _dedupe_related_events(events: List[ScheduleEventSchema]) -> List[ScheduleEventSchema]:
    deduped: list[ScheduleEventSchema] = []
    seen: set[tuple[str, int]] = set()
    for event in events:
        key = _build_event_key(event)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(event)
    return deduped


def _is_movable_personal_scheduled_event(event: ScheduleEventSchema) -> bool:
    return (
        event.source == "personal_todo"
        and event.kind == "scheduled_todo"
        and event.is_blocking
        and event.is_flexible
        and _has_explicit_range(event)
        and event.status not in {"done", "cancelled"}
    )


def _get_blocking_events(
    schedule_events: List[ScheduleEventSchema],
    exclude_event_keys: set[tuple[str, int]] | None = None,
) -> List[ScheduleEventSchema]:
    excluded = exclude_event_keys or set()
    return [
        event for event in schedule_events
        if event.is_blocking
        and _has_explicit_range(event)
        and _build_event_key(event) not in excluded
    ]


def _build_day_window(anchor: datetime, start_hour: int = SUGGESTION_DAY_START_HOUR, end_hour: int = SUGGESTION_DAY_END_HOUR) -> tuple[datetime, datetime]:
    day_start = datetime.combine(anchor.date(), time(start_hour), tzinfo=anchor.tzinfo)
    day_end = datetime.combine(anchor.date(), time(end_hour), tzinfo=anchor.tzinfo)
    return day_start, day_end


def _collect_overlapping_blocker_ids(
    start_time: datetime,
    end_time: datetime,
    blocking_events: List[ScheduleEventSchema],
) -> List[int]:
    blocker_ids: list[int] = []
    for event in blocking_events:
        if event.start_time < end_time and start_time < event.end_time:
            blocker_ids.append(event.source_id)
    return sorted(set(blocker_ids))


def _find_free_windows(
    blocking_events: List[ScheduleEventSchema],
    window_start: datetime,
    window_end: datetime,
) -> List[tuple[datetime, datetime]]:
    if window_start >= window_end:
        return []

    relevant_events = sorted(
        [
            event for event in blocking_events
            if event.start_time < window_end and window_start < event.end_time
        ],
        key=lambda item: (item.start_time, item.end_time, item.source, item.source_id),
    )
    free_windows: list[tuple[datetime, datetime]] = []
    cursor = window_start
    for event in relevant_events:
        occupied_start = max(window_start, event.start_time)
        occupied_end = min(window_end, event.end_time)
        if cursor < occupied_start:
            free_windows.append((cursor, occupied_start))
        if occupied_end > cursor:
            cursor = occupied_end
    if cursor < window_end:
        free_windows.append((cursor, window_end))
    return free_windows


def _merge_time_intervals(intervals: List[tuple[datetime, datetime]]) -> List[tuple[datetime, datetime]]:
    if not intervals:
        return []
    merged: list[tuple[datetime, datetime]] = []
    for start_time, end_time in sorted(intervals, key=lambda item: (item[0], item[1])):
        if not merged or start_time >= merged[-1][1]:
            merged.append((start_time, end_time))
            continue
        merged[-1] = (merged[-1][0], max(merged[-1][1], end_time))
    return merged


def _calculate_interval_union_hours(events: List[ScheduleEventSchema]) -> float:
    merged_intervals = _merge_time_intervals(
        [(event.start_time, event.end_time) for event in events if _has_explicit_range(event)]
    )
    return sum(
        max(0.0, (end_time - start_time).total_seconds() / 3600)
        for start_time, end_time in merged_intervals
    )


def _build_candidate_reschedule_options(
    target_event: ScheduleEventSchema,
    related_blockers: List[ScheduleEventSchema],
    schedule_events: List[ScheduleEventSchema],
    max_candidates: int = 3,
) -> List[ScheduleCandidateRescheduleOptionSchema]:
    if not _is_movable_personal_scheduled_event(target_event):
        return []

    duration = target_event.end_time - target_event.start_time
    blocking_events = _get_blocking_events(
        schedule_events,
        exclude_event_keys={_build_event_key(target_event)},
    )
    day_offsets = [0, -1, 1]
    candidates: list[ScheduleCandidateRescheduleOptionSchema] = []
    seen_slots: set[tuple[datetime, datetime]] = set()

    for day_offset in day_offsets:
        anchor = target_event.start_time + timedelta(days=day_offset)
        window_start, window_end = _build_day_window(anchor)
        free_windows = _find_free_windows(blocking_events, window_start, window_end)
        for free_start, free_end in free_windows:
            if free_end - free_start < duration:
                continue
            candidate_start = free_start
            candidate_end = candidate_start + duration
            slot_key = (candidate_start, candidate_end)
            if slot_key in seen_slots:
                continue
            seen_slots.add(slot_key)
            candidates.append(
                ScheduleCandidateRescheduleOptionSchema(
                    action="reschedule_personal_todo",
                    target_event_id=target_event.source_id,
                    suggested_start_time=candidate_start,
                    suggested_end_time=candidate_end,
                    reason="避开当前固定阻塞时间段后的可用时间块",
                    blocked_by_event_ids=sorted({_id for event in related_blockers for _id in [event.source_id]}),
                    rank=len(candidates) + 1,
                    confidence="high" if day_offset == 0 else "medium",
                )
            )
            if len(candidates) >= max_candidates:
                return candidates

    return candidates


def _build_candidate_focus_blocks(
    deadline_events: List[ScheduleEventSchema],
    pileup: SchedulePileupSchema,
    schedule_events: List[ScheduleEventSchema],
    max_candidates: int = 3,
) -> List[ScheduleCandidateFocusBlockSchema]:
    if not deadline_events:
        return []

    earliest_deadline = min(event.deadline for event in deadline_events)
    target_duration = timedelta(minutes=DEFAULT_FOCUS_BLOCK_MINUTES)
    blocking_events = _get_blocking_events(schedule_events)
    candidates: list[ScheduleCandidateFocusBlockSchema] = []
    seen_slots: set[tuple[datetime, datetime]] = set()
    day_offsets = [0, -1]

    for day_offset in day_offsets:
        anchor = earliest_deadline + timedelta(days=day_offset)
        window_start, window_end = _build_day_window(anchor)
        window_end = min(window_end, earliest_deadline)
        if window_end <= window_start:
            continue
        free_windows = _find_free_windows(blocking_events, window_start, window_end)
        for free_start, free_end in sorted(free_windows, key=lambda item: item[1], reverse=True):
            candidate_end = free_end
            candidate_start = max(free_start, candidate_end - target_duration)
            if candidate_end - candidate_start < timedelta(hours=1):
                continue
            if candidate_start.hour < SUGGESTION_DAY_START_HOUR:
                continue
            slot_key = (candidate_start, candidate_end)
            if slot_key in seen_slots:
                continue
            seen_slots.add(slot_key)
            candidates.append(
                ScheduleCandidateFocusBlockSchema(
                    suggested_start_time=candidate_start,
                    suggested_end_time=candidate_end,
                    intended_for_event_ids=sorted({event.source_id for event in deadline_events}),
                    reason="截止前的可用专注时间块，不与现有阻塞安排重叠",
                    rank=len(candidates) + 1,
                    confidence="high" if day_offset == 0 else "medium",
                )
            )
            if len(candidates) >= max_candidates:
                return candidates

    return candidates


def _build_candidate_split_options(
    target_event: ScheduleEventSchema,
    schedule_events: List[ScheduleEventSchema],
    max_candidates: int = 1,
) -> List[ScheduleCandidateSplitOptionSchema]:
    if not _is_movable_personal_scheduled_event(target_event):
        return []

    total_duration = target_event.end_time - target_event.start_time
    if total_duration < timedelta(hours=4):
        return []

    blocking_events = _get_blocking_events(
        schedule_events,
        exclude_event_keys={_build_event_key(target_event)},
    )
    total_duration_minutes = int(total_duration.total_seconds() // 60)
    candidates: list[ScheduleCandidateSplitOptionSchema] = []
    day_offsets = [0, -1, 1]

    segments: list[ScheduleCandidateSplitSegmentSchema] = []
    remaining_duration = total_duration
    for day_offset in day_offsets:
        anchor = target_event.start_time + timedelta(days=day_offset)
        window_start, window_end = _build_day_window(anchor)
        free_windows = _find_free_windows(blocking_events, window_start, window_end)
        for free_start, free_end in free_windows:
            if len(segments) >= 3:
                break
            available_duration = free_end - free_start
            if available_duration < timedelta(minutes=MIN_SPLIT_SEGMENT_MINUTES):
                continue
            segment_end = free_start + min(available_duration, remaining_duration)
            if segment_end - free_start < timedelta(minutes=MIN_SPLIT_SEGMENT_MINUTES):
                continue
            segments.append(
                ScheduleCandidateSplitSegmentSchema(
                    suggested_start_time=free_start,
                    suggested_end_time=segment_end,
                )
            )
            remaining_duration -= segment_end - free_start
            if remaining_duration <= timedelta():
                candidates.append(
                    ScheduleCandidateSplitOptionSchema(
                        action="split_personal_todo",
                        target_event_id=target_event.source_id,
                        segments=segments,
                        total_duration_minutes=total_duration_minutes,
                        reason="避开课程与固定阻塞时间段，将长任务拆分为更可执行的多个时间块",
                        rank=1,
                        confidence="medium",
                    )
                )
                return candidates[:max_candidates]

    return []


def detect_daily_scheduled_overload(
    schedule_events: List[ScheduleEventSchema],
    daily_hours_threshold: int = 8,
    daily_item_threshold: int = 6,
) -> List[SchedulePileupSchema]:
    """
    V1 daily overload uses the union of blocking explicit start/end ranges by day.
    Deadline-only tasks and official background markers do not occupy scheduled hours.
    When a background marker is active, the daily threshold is reduced slightly to model
    background pressure without turning the marker itself into scheduled load.
    """
    ranged_events = [
        event for event in schedule_events
        if event.is_blocking and _has_explicit_range(event)
    ]
    background_markers = [
        event for event in schedule_events
        if _is_background_pressure_marker(event) and _has_explicit_range(event)
    ]
    events_by_day: dict[date, list[ScheduleEventSchema]] = {}
    for event in ranged_events:
        events_by_day.setdefault(event.start_time.date(), []).append(event)

    pileups: list[SchedulePileupSchema] = []
    for day, events in sorted(events_by_day.items(), key=lambda item: item[0]):
        day_start = datetime.combine(day, time.min, tzinfo=events[0].start_time.tzinfo)
        day_end = day_start + timedelta(days=1)
        has_background_pressure = any(
            _time_window_overlaps(marker, day_start, day_end)
            for marker in background_markers
        )
        effective_hours_threshold = daily_hours_threshold * (0.9 if has_background_pressure else 1.0)
        effective_item_threshold = daily_item_threshold - 1 if has_background_pressure else daily_item_threshold
        occupied_hours = _calculate_interval_union_hours(events)
        if occupied_hours < effective_hours_threshold and len(events) < effective_item_threshold:
            continue

        reason = f"当日显式阻塞时间段合并后占用约 {round(occupied_hours, 1)} 小时，安排偏满"
        if has_background_pressure:
            reason += "；当前处于官方背景提醒周期，建议预留缓冲"
        pileups.append(
            SchedulePileupSchema(
                start_time=min(event.start_time for event in events),
                end_time=max(event.end_time for event in events),
                event_count=len(events),
                level="high" if occupied_hours >= effective_hours_threshold + 2 else "medium",
                reason=reason,
                events=sorted(events, key=_build_schedule_sort_key),
            )
        )

    return pileups


def build_schedule_suggestions(
    schedule_events: List[ScheduleEventSchema],
    conflicts: List[ScheduleConflictSchema],
    pileups: List[SchedulePileupSchema],
) -> List[ScheduleSuggestionSchema]:
    """
    V1 suggestions are deterministic and conservative:
    move only flexible personal tasks, never official fixed events, and for risky deadlines
    suggest creating an earlier focus block instead of inventing a reschedule target.
    """
    suggestions: list[ScheduleSuggestionSchema] = []
    suggestion_keys: set[tuple] = set()
    background_markers = [event for event in schedule_events if _is_background_pressure_marker(event)]
    deadline_events_by_day: dict[date, list[ScheduleEventSchema]] = {}
    conflicts_by_personal_event: dict[int, list[ScheduleConflictSchema]] = {}
    for event in schedule_events:
        if _is_pending_deadline_todo(event):
            deadline_events_by_day.setdefault(event.deadline.date(), []).append(event)
    for conflict in conflicts:
        for event in (conflict.left_event, conflict.right_event):
            if _is_movable_personal_scheduled_event(event):
                conflicts_by_personal_event.setdefault(event.source_id, []).append(conflict)

    for target_event_id, target_conflicts in sorted(conflicts_by_personal_event.items()):
        all_events = [
            event
            for conflict in target_conflicts
            for event in (conflict.left_event, conflict.right_event)
        ]
        target_event = next((event for event in all_events if event.source == "personal_todo" and event.source_id == target_event_id), None)
        if target_event is None:
            continue
        blocking_related_events = [
            event for event in all_events
            if _build_event_key(event) != _build_event_key(target_event)
        ]
        candidate_reschedule_options = _build_candidate_reschedule_options(
            target_event=target_event,
            related_blockers=blocking_related_events,
            schedule_events=schedule_events,
        )
        candidate_split_options = _build_candidate_split_options(
            target_event=target_event,
            schedule_events=schedule_events,
        )
        key = ("review_personal_conflict", target_event.source_id)
        if key in suggestion_keys:
            continue
        suggestion_keys.add(key)
        primary_candidate = candidate_reschedule_options[0] if candidate_reschedule_options else None
        suggestions.append(
            ScheduleSuggestionSchema(
                type="review_personal_task_time",
                level="high",
                message="个人事项与固定时间段冲突，建议调整个人事项时间块",
                related_events=_dedupe_related_events([target_event, *blocking_related_events]),
                suggested_start_time=primary_candidate.suggested_start_time if primary_candidate else None,
                suggested_end_time=primary_candidate.suggested_end_time if primary_candidate else None,
                candidate_reschedule_options=candidate_reschedule_options,
                candidate_split_options=candidate_split_options,
            )
        )

    for pileup in pileups:
        has_background_pressure = any(
            _time_window_overlaps(marker, pileup.start_time, pileup.end_time)
            for marker in background_markers
        )
        deadline_events = [event for event in pileup.events if _is_pending_deadline_todo(event)]
        if not deadline_events and pileup.start_time is not None:
            deadline_events = deadline_events_by_day.get(pileup.start_time.date(), [])
        if deadline_events:
            earliest_deadline = min(event.deadline for event in deadline_events)
            key = ("focus_before_deadline", earliest_deadline, tuple(event.source_id for event in deadline_events))
            if key not in suggestion_keys:
                suggestion_keys.add(key)
                has_course_blockers = any(event.source == "course_schedule_event" for event in pileup.events)
                candidate_focus_blocks = _build_candidate_focus_blocks(deadline_events, pileup, schedule_events)
                primary_candidate = candidate_focus_blocks[0] if candidate_focus_blocks else None
                suggestions.append(
                    ScheduleSuggestionSchema(
                        type="create_focus_block",
                        level=pileup.level,
                        message="当天固定课程等时间块较满，建议在剩余空档提前创建专注时间块处理截止事项"
                        if has_course_blockers
                        else (
                            "截止事项偏密集，建议提前创建专注时间块处理"
                            if not has_background_pressure
                            else "当前处于官方背景提醒周期，建议更早规划截止事项处理时间"
                        ),
                        related_events=_dedupe_related_events([*deadline_events, *pileup.events]),
                        suggested_start_time=primary_candidate.suggested_start_time if primary_candidate else None,
                        suggested_end_time=primary_candidate.suggested_end_time if primary_candidate else None,
                        suggested_deadline=earliest_deadline,
                        candidate_focus_blocks=candidate_focus_blocks,
                    )
                )
            continue

        flexible_personal_events = [
            event for event in pileup.events
            if event.source == "personal_todo" and event.is_flexible
        ]
        if not flexible_personal_events:
            continue
        candidate = sorted(flexible_personal_events, key=lambda item: (item.priority or 0, item.source_id))[0]
        key = ("rebalance_day", candidate.source_id, pileup.start_time.date() if pileup.start_time else None)
        if key in suggestion_keys:
            continue
        suggestion_keys.add(key)
        candidate_split_options = _build_candidate_split_options(candidate, schedule_events)
        if candidate_split_options:
            suggestions.append(
                ScheduleSuggestionSchema(
                    type="split_personal_task",
                    level=pileup.level,
                    message="该个人事项时长较长，建议拆分为多个可执行时间块",
                    related_events=_dedupe_related_events([candidate, *pileup.events]),
                    candidate_split_options=candidate_split_options,
                )
            )
            continue
        suggestions.append(
            ScheduleSuggestionSchema(
                type="rebalance_personal_schedule",
                level=pileup.level,
                message="当日显式时间段过于集中，优先调整个人事项或拆分执行"
                if not has_background_pressure
                else "当前处于官方背景提醒周期，避免在这一时段继续堆叠过多个人安排",
                related_events=[candidate],
                suggested_start_time=pileup.end_time,
            )
        )

    return suggestions


def build_schedule_analysis_summary(
    schedule_events: List[ScheduleEventSchema],
    conflicts: List[ScheduleConflictSchema],
    pileups: List[SchedulePileupSchema],
    suggestions: List[ScheduleSuggestionSchema],
) -> ScheduleAnalysisSummarySchema:
    return ScheduleAnalysisSummarySchema(
        total_events=len(schedule_events),
        blocking_event_count=sum(1 for event in schedule_events if event.is_blocking),
        flexible_event_count=sum(1 for event in schedule_events if event.is_flexible),
        deadline_only_count=sum(1 for event in schedule_events if event.deadline is not None and not _has_explicit_range(event)),
        conflict_count=len(conflicts),
        pileup_count=len(pileups),
        suggestion_count=len(suggestions),
        risk_level=_calculate_risk_level(conflicts, pileups, suggestions),
    )


def build_schedule_ai_context(
    summary: ScheduleAnalysisSummarySchema,
    conflicts: List[ScheduleConflictSchema],
    pileups: List[SchedulePileupSchema],
    suggestions: List[ScheduleSuggestionSchema],
) -> ScheduleAiContextSchema:
    top_conflicts = [
        {
            "type": conflict.conflict_type,
            "level": conflict.level,
            "reason": conflict.reason,
            "left_event_id": conflict.left_event.source_id,
            "left_event_title": conflict.left_event.title,
            "right_event_id": conflict.right_event.source_id,
            "right_event_title": conflict.right_event.title,
        }
        for conflict in conflicts[:3]
    ]
    top_pileups = [
        {
            "level": pileup.level,
            "reason": pileup.reason,
            "event_count": pileup.event_count,
            "event_ids": [event.source_id for event in pileup.events[:5]],
        }
        for pileup in pileups[:3]
    ]
    top_candidate_actions: list[dict[str, Any]] = []
    for suggestion in suggestions:
        for option in suggestion.candidate_reschedule_options[:2]:
            top_candidate_actions.append(
                {
                    "type": "reschedule_option",
                    "target_event_id": option.target_event_id,
                    "suggested_start_time": option.suggested_start_time,
                    "suggested_end_time": option.suggested_end_time,
                    "reason": option.reason,
                    "rank": option.rank,
                }
            )
        for option in suggestion.candidate_focus_blocks[:2]:
            top_candidate_actions.append(
                {
                    "type": "focus_block_option",
                    "intended_for_event_ids": option.intended_for_event_ids,
                    "suggested_start_time": option.suggested_start_time,
                    "suggested_end_time": option.suggested_end_time,
                    "reason": option.reason,
                    "rank": option.rank,
                }
            )
        for option in suggestion.candidate_split_options[:1]:
            top_candidate_actions.append(
                {
                    "type": "split_option",
                    "target_event_id": option.target_event_id,
                    "segment_count": len(option.segments),
                    "total_duration_minutes": option.total_duration_minutes,
                    "reason": option.reason,
                    "rank": option.rank,
                }
            )
        if len(top_candidate_actions) >= 5:
            break

    return ScheduleAiContextSchema(
        schedule_risk_summary=f"risk={summary.risk_level}; conflicts={summary.conflict_count}; pileups={summary.pileup_count}; suggestions={summary.suggestion_count}",
        top_conflicts=top_conflicts,
        top_pileups=top_pileups,
        top_candidate_actions=top_candidate_actions[:5],
    )


def analyze_schedule_timeline(
    schedule_events: List[ScheduleEventSchema],
) -> dict[str, Any]:
    """
    V1 analysis is read-only and deterministic.
    It only uses explicit time ranges, stored deadlines, and current status fields; it does not estimate missing durations.
    """
    sorted_events = sorted(schedule_events, key=_build_schedule_sort_key)
    conflicts = detect_time_overlap_conflicts(sorted_events)
    pileups = detect_deadline_clusters(sorted_events)
    pileups.extend(detect_daily_scheduled_overload(sorted_events))
    suggestions = build_schedule_suggestions(sorted_events, conflicts, pileups)
    summary = build_schedule_analysis_summary(sorted_events, conflicts, pileups, suggestions)
    ai_context = build_schedule_ai_context(summary, conflicts, pileups, suggestions)
    return {
        "summary": summary,
        "conflicts": conflicts,
        "pileups": pileups,
        "suggestions": suggestions,
        "ai_context": ai_context,
    }


async def build_schedule_aggregate_data(
    db: AsyncSession,
    user_id: int,
    start_time: datetime | None = None,
    end_time: datetime | None = None,
    include_blackboard: bool = True,
    include_background_markers: bool = True,
    include_completed: bool = True,
    include_course_schedule: bool = True,
) -> ScheduleAggregateData:
    personal_todos = await load_personal_todos_by_window(
        db=db,
        user_id=user_id,
        start_time=start_time,
        end_time=end_time,
        include_blackboard=include_blackboard,
        include_completed=include_completed,
    )
    official_events = await load_official_calendar_events_by_window(
        db=db,
        start_time=start_time,
        end_time=end_time,
        include_background_markers=include_background_markers,
    )
    course_events = []
    if include_course_schedule:
        course_events = await list_course_schedule_events(
            db=db,
            user_id=user_id,
            start_time=start_time,
            end_time=end_time,
        )
    schedule_events = merge_schedule_timeline(personal_todos, official_events, course_events)
    analysis = analyze_schedule_timeline(schedule_events)
    return ScheduleAggregateData(
        todos=personal_todos,
        official_events=official_events,
        schedule_events=schedule_events,
        conflicts=analysis["conflicts"],
        pileups=analysis["pileups"],
        suggestions=analysis["suggestions"],
        summary=analysis["summary"],
        ai_context=analysis["ai_context"],
    )
